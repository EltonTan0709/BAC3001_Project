"""Read model inputs from a recalculated XLSX export; never evaluate Excel formulas.

Cell mappings below deliberately target the submitted workbook's six-period explicit
forecast and five-year transition. Missing caches or a changed layout fail loudly.
"""
from datetime import date, datetime
import hashlib
import math
import os
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / "data" / "Business_Valuation.xlsx"


class WorkbookInputs:
    def __init__(self, path=None):
        self.path = Path(path or os.environ.get("BAC3001_WORKBOOK", DEFAULT_WORKBOOK)).resolve()
        if not self.path.is_file():
            raise FileNotFoundError(f"Workbook not found: {self.path}")
        self.sha256 = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.book = load_workbook(self.path, data_only=True, read_only=False)
        required = {"Valuation", "Assumptions", "Historical base", "Financial assets",
                    "Working capital", "Capital and FCFF", "Transition", "Simulation inputs"}
        missing = required.difference(self.book.sheetnames)
        if missing:
            raise ValueError(f"Missing workbook tabs: {sorted(missing)}")
        expected = ["H2 2026", 2027, 2028, 2029, 2030, 2031]
        actual = [self.book["Assumptions"].cell(6, c).value for c in range(2, 8)]
        if [str(v) for v in actual] != [str(v) for v in expected]:
            raise ValueError("Expected Assumptions!B6:G6 = H2 2026, 2027 ... 2031. Update the mapping for a new horizon.")
        if self.number("Assumptions", "B147") != 5:
            raise ValueError("This replica implements a five-year transition (2032-2036).")
        if "capacity benchmark" not in str(self.book["Assumptions"]["A192"].value).lower():
            raise ValueError("This engine requires the post-build revision: Assumptions!A192 "
                             "must identify capital intensity relative to the capacity benchmark.")

    def number(self, sheet, cell):
        value = self.book[sheet][cell].value
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError(f"{sheet}!{cell} has no finite numeric cache ({value!r}). "
                             "Recalculate and save/export the workbook in Excel or Google Sheets first.")
        return float(value)

    def row(self, sheet, row, first=2, last=7):
        from openpyxl.utils import get_column_letter
        return np.array([self.number(sheet, f"{get_column_letter(c)}{row}")
                         for c in range(first, last + 1)])

    def common_path(self, first_row):
        base = self.row("Assumptions", first_row)
        for row in (first_row + 1, first_row + 2):
            if not np.allclose(base, self.row("Assumptions", row), rtol=0, atol=1e-10):
                raise ValueError(f"Assumptions rows {first_row}:{first_row + 2} now differ by scenario. "
                                 "Extend the Python mapping before relying on simulation results.")
        return base

    def scalar_cases(self, first_row):
        values = []
        for row in range(first_row, first_row + 3):
            path = self.row("Assumptions", row)
            if not np.allclose(path, path[0], rtol=0, atol=1e-10):
                raise ValueError(f"Assumptions row {row} is no longer constant across years; extend the scalar mapping.")
            values.append(path[0])
        return tuple(values)  # Base, Bull, Bear

    def inputs(self):
        n = self.number
        a = lambda cell: n("Assumptions", cell)
        history_rows = {"search": 8, "yt": 9, "net": 10, "subs": 11, "cloud": 13, "ob": 14}
        hist = lambda col: {k: n("Historical base", f"{col}{r}") for k, r in history_rows.items()}
        fy25, h125, h126 = hist("C"), hist("D"), hist("E")
        rows = {"search": 8, "yt": 13, "net": 18, "subs": 23, "cloud_demand": 28,
                "ob_rev": 33, "cloud_margin": 43, "ob_cost": 48, "capex_int": 58,
                "inventory": 109, "tac_change": 154}
        scen = {}
        for key, start in rows.items():
            scen[key] = {}
            for case, row in zip(("base", "bull", "bear"), range(start, start + 3)):
                if key == "capex_int":
                    values = np.r_[np.nan, self.row("Assumptions", row, 3)]
                else:
                    values = self.row("Assumptions", row)
                scen[key][case] = values
        scalar = {"commission": self.scalar_cases(63), "hw_life": self.scalar_cases(68),
                  "research_elast": self.scalar_cases(53)}
        for key, start in {"fund_defer": 178, "fund_pressure": 183, "capital_intensity": 193,
                           "term_growth": 198, "ai_unit": 208, "ai_workload": 213,
                           "research_h2": 223}.items():
            scalar[key] = tuple(a(f"B{r}") for r in range(start, start + 3))
        # These components are common across the three cases in the source workbook.
        # Receivables use days / 365; liabilities offset operating assets.
        nowc_ex_inv = self.common_path(104) / 365 + self.common_path(114)
        for first_row in (119, 124, 129, 134, 139):
            nowc_ex_inv -= self.common_path(first_row)
        cip, ppe = n("Historical base", "E74"), n("Historical base", "E72")
        opening_hw = (ppe - cip) * a("B89")
        raw_date = self.book["Valuation"]["B4"].value
        val_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date
        if not isinstance(val_date, date):
            raise ValueError("Valuation!B4 must contain an Excel date.")
        tail_years = int(a("B238"))
        if tail_years < 1 or tail_years > 5:
            raise ValueError("Intangible runoff must finish within the five-year transition.")
        tail = np.r_[np.full(tail_years, a("B237") / tail_years), np.zeros(6 - tail_years)]
        margin_override = self.book["Assumptions"]["B148"].value
        if margin_override is not None:
            margin_override = a("B148")
        periods = self.row("Capital and FCFF", 7)
        if not np.array_equal(periods, [0.5, 1, 1, 1, 1, 1]):
            raise ValueError("Explicit period lengths must be 0.5, 1, 1, 1, 1, 1.")
        for row in (173, 188):
            vals = [a(f"B{r}") for r in range(row, row + 3)]
            if not np.allclose(vals, vals[0]):
                raise ValueError(f"Assumptions rows {row}:{row+2} now differ; extend the common-exposure mapping.")
        return dict(
            FY25=fy25, H1_25=h125, H1_26=h126,
            H2_25={k: fy25[k] - h125[k] for k in fy25},
            H1_26_REV=n("Historical base", "E16"),
            OB_LOSS_FY25=-n("Historical base", "C22"),
            OB_LOSS_H1_25=-n("Historical base", "D22"),
            OB_LOSS_H1_26=-n("Historical base", "E22"),
            ALPHA_H1_26=-n("Historical base", "E23"), REF_DA_RATIO=a("B95"),
            REF_NET_SHARE=h126["net"] / sum(h126[k] for k in ("search", "yt", "net")),
            SCEN=scen, SCALAR=scalar, SERV_MARGIN=self.common_path(38),
            DRIVER_OF={"research_h2": "research_elast", "fund_defer": "funding",
                       "fund_pressure": "funding", "ob_rev": "ob_cost"},
            SERV_TAC_DIFF=a("B227"), AI_EXPOSURE=a("B217"),
            RESEARCH_SHARE=a("B220"), OTHER_SHARED_GROWTH=a("B221"),
            INTANG=self.row("Capital and FCFF", 25), INTANG_TAIL=tail,
            INTANG_TAX_DEDUCTIBLE=a("B239"), NOWC_EX_INV=nowc_ex_inv,
            NOWC_OPEN=n("Working capital", "C29"),
            OPERATING_CASH_RATIO=n("Financial assets", "B30"),
            OPENING_TTM_REVENUE=n("Historical base", "F16"),
            BACKLOG=n("Cloud demand", "B7"), BACKLOG_SHARE=self.row("Assumptions", 167),
            MIN_NEW_SHARE=a("B171"), FUND_EXPOSURE=a("B173"), DELIVERY_SENSITIVE=a("B188"),
            CAPACITY_REQ=self.row("Assumptions", 234), HIST_COMMISSION=a("B170"),
            CAPEX_H2=a("B93"), PPE_OPEN=ppe, CIP_OPEN=cip,
            HW_OPEN=opening_hw, OTH_OPEN=ppe-cip-opening_hw,
            HW_OPEN_LIFE=a("B90"), OTH_OPEN_LIFE=a("B91"),
            HW_SHARE=float(self.common_path(73)[0]), OTH_LIFE=a("B92"), PERIOD=periods,
            WACC=n("Valuation", "B29"), TAX_TERMINAL=a("B149"),
            TAX_EXPLICIT=self.common_path(78), VAL_DATE=val_date,
            MID_YEAR=bool(n("Valuation", "B24")), TERM_MARGIN_OVERRIDE=margin_override,
            STUB_FRAC=n("Valuation", "B36"),
            YRS_EX=np.array([(date(2026+i, 12, 31)-val_date).days/365 for i in range(6)]),
            YRS_TR=np.array([(date(2032+i, 12, 31)-val_date).days/365 for i in range(5)]),
        )


def load_inputs(path=None):
    return WorkbookInputs(path)
