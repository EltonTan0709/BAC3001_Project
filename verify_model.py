"""Reconcile Python cash flows and valuation to the included workbook's Base caches."""
import argparse
import json
from pathlib import Path

import numpy as np
import alphabet_model as am


def verify(output_dir=Path("results")):
    w = am.WORKBOOK
    if w.book["Valuation"]["B3"].value != "Base":
        raise ValueError("Select Base, recalculate and save/export the workbook before verification.")
    b = am.value()
    checks = {}

    def compare(name, actual, expected, tolerance=0.02, unit="usdm"):
        error = float(np.max(np.abs(np.asarray(actual) - np.asarray(expected))))
        checks[name] = {f"maximum_absolute_error_{unit}": error, f"tolerance_{unit}": tolerance,
                        "passed": bool(error <= tolerance)}

    n = w.number
    compare("base_operating_ev", b["ev"], n("Valuation", "B48"))
    compare("explicit_capex", b["capex"], w.row("Capital and FCFF", 12))
    compare("explicit_depreciation", b["dep"], w.row("Capital and FCFF", 24))
    compare("explicit_nopat", b["nopat"], w.row("Capital and FCFF", 31))
    compare("explicit_fcff", b["fcff"], w.row("Valuation", 37))
    compare("explicit_operating_wc", b["operating_wc"], w.row("Working capital", 59, 5, 10))
    compare("explicit_wc_investment", b["d_operating_wc"], w.row("Working capital", 62, 5, 10))
    compare("transition_fcff", b["transition_fcff"], w.row("Transition", 16, 3, 7))
    compare("transition_capex", b["transition_capex"], w.row("Capital and FCFF", 12, 8, 12))
    compare("transition_depreciation", b["transition_dep"], w.row("Capital and FCFF", 24, 8, 12))
    compare("transition_cip", b["transition_cip"], w.row("Capital and FCFF", 15, 8, 12))
    compare("transition_ppe", b["transition_ppe"], w.row("Capital and FCFF", 26, 8, 12))
    compare("transition_roc", b["transition_roc"], w.row("Capital and FCFF", 40, 8, 12), tolerance=1e-8, unit="ratio")
    compare("pv_explicit", b["pv_ex"], sum(w.row("Valuation", 39)))
    compare("pv_transition", b["pv_tr"], sum(w.row("Transition", 20, 3, 7)))
    compare("pv_terminal", b["pv_tv"], n("Valuation", "B46"))
    cases = {}
    for case, row in (("bear", 46), ("base", 47), ("bull", 48)):
        r = am.scenario(case)
        cases[case] = {"operating_ev_usdm": float(r["ev"]), "starting_return": float(r["r0"]),
                       "terminal_return": float(r["RT"])}
        compare(f"{case}_scenario_snapshot", r["ev"] / 1000,
                n("Simulation inputs", f"B{row}"), tolerance=0.1)
        # These three snapshot comparisons are in USD billions, not millions.
        checks[f"{case}_scenario_snapshot"]["maximum_absolute_error_usdbn"] = checks[f"{case}_scenario_snapshot"].pop("maximum_absolute_error_usdm")
        checks[f"{case}_scenario_snapshot"]["tolerance_usdbn"] = checks[f"{case}_scenario_snapshot"].pop("tolerance_usdm")
        checks[f"{case}_capital_integrity"] = {"passed": bool(
            np.all(r["transition_cip"] >= 0) and np.all(r["transition_capex"] >= 0)
            and np.allclose(r["transition_ppe"]-r["transition_cip"], r["transition_target"], atol=.02)
            and r["RT"] >= am.WACC)}
    # Behaviour checks distinguish sensible cash-flow responses from mere replica agreement.
    slow = am.value(overrides={"commission": .4})
    fast = am.value(overrides={"commission": .8})
    short = am.value(overrides={"hw_life": 4})
    long = am.value(overrides={"hw_life": 6})
    checks["slower_commissioning_reduces_value"] = {"passed": bool(slow["ev"] < fast["ev"])}
    checks["longer_life_reduces_postbuild_replacement"] = {"passed": bool(
        long["transition_capex"].sum() < short["transition_capex"].sum())}
    bridge = n("Valuation", "B49")-n("Valuation", "B50")
    w0 = am.WACC
    try:
        grid=[]
        for row in range(17, 22):
            from openpyxl.utils import get_column_letter
            g=n("Sensitivity", f"A{row}")
            values=[]
            for col in range(2, 7):
                letter=get_column_letter(col)
                am.WACC=n("Sensitivity", f"{letter}16")
                values.append((am.value(overrides={"term_growth": g})["ev"]+bridge)/1e6)
            grid.append(values)
        expected=[w.row("Sensitivity", row, 2, 6) for row in range(17,22)]
        compare("all_25_sensitivity_cells", grid, expected, tolerance=1e-8, unit="usdtn")
    finally:
        am.WACC=w0
    result = {"workbook": w.path.name, "workbook_sha256": w.sha256,
              "valuation_date": am.VAL_DATE.isoformat(), "checks": checks,
              "deterministic_scenarios": cases,
              "passed": all(c["passed"] for c in checks.values())}
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    (Path(output_dir) / "verification.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))
    if not result["passed"]:
        raise SystemExit("Reconciliation failed. Check workbook caches, scenario snapshot and Python mappings before using outputs.")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("results"))
    args = parser.parse_args()
    if args.workbook:
        am.load_workbook(args.workbook)
    verify(args.output_dir)
